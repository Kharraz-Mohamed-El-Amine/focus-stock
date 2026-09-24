import re
import datetime
import logging
from typing import Dict, Any, List, Optional
import openpyxl

from stock.models import Site, Reference, Mouvement, Trajet, TrajetMouvement

logger = logging.getLogger(__name__)


def parse_date_value(val: Any) -> Optional[datetime.date]:
    """Parse various date formats from Excel (datetime object, string dd/mm/yyyy or yyyy-mm-dd)."""
    if val is None:
        return None
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.date() if isinstance(val, datetime.datetime) else val
    if isinstance(val, str):
        val_str = val.strip()
        m = re.search(r'(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})', val_str)
        if m:
            try:
                return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            except ValueError:
                pass
        m2 = re.search(r'(\d{4})[/.-](\d{1,2})[/.-](\d{1,2})', val_str)
        if m2:
            try:
                return datetime.date(int(m2.group(1)), int(m2.group(2)), int(m2.group(3)))
            except ValueError:
                pass
    return None


def parse_int_value(val: Any) -> Optional[int]:
    """Safely convert numerical/float/string values to integer."""
    if val is None:
        return None
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return None


def clean_pn_value(val: Any) -> Optional[str]:
    """Clean Part Number string."""
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() in ('none', 'without reception', 'total quantity'):
        return None
    return s


PLANT_MAPPING = {
    "TOUBQAL": "TOUBKAL",
    "TOUBKAL": "TOUBKAL",
    "TFZ": "TFZ",
    "MTST": "MTST",
}


def normalize_plant_code(val: Any, default: Optional[str] = None) -> Optional[str]:
    """
    Normalise la valeur Plant :
    1. strip() des espaces
    2. mise en majuscules (upper())
    3. correspondance via le dictionnaire PLANT_MAPPING
    """
    if not val:
        return default
    val_str = str(val).strip()
    val_upper = val_str.upper()

    if val_upper in PLANT_MAPPING:
        return PLANT_MAPPING[val_upper]

    # Vérification des sous-chaînes si le libellé contient le nom (ex: 'TOUBKAL-BLOCKED STOCK')
    for k, v in PLANT_MAPPING.items():
        if k in val_upper:
            return v

    return default or val_upper[:20]


def get_or_create_site_with_warning(
    site_code: str,
    row_num: Optional[int] = None,
    avertissements: Optional[List[str]] = None,
    erreurs: Optional[List[str]] = None
) -> Optional[Site]:
    """
    Récupère ou crée le Site. Si le code normalisé n'existe dans NI la base de données
    NI le dictionnaire PLANT_MAPPING, NE PAS créer le Site automatiquement.
    À la place, log un avertissement clair avec le numéro de ligne et le code rencontré,
    ajoute une erreur au rapport et retourne None pour que la ligne soit ignorée.
    """
    site = Site.objects.filter(code=site_code).first()
    if site:
        return site

    known_in_mapping = (site_code in PLANT_MAPPING) or (site_code in PLANT_MAPPING.values())
    if not known_in_mapping:
        row_str = f"Ligne {row_num} : " if row_num else ""
        msg = f"{row_str}Avertissement : Le site '{site_code}' n'existe ni en base de données ni dans PLANT_MAPPING. Ligne ignorée."
        logger.warning(msg)
        if avertissements is not None and msg not in avertissements:
            avertissements.append(msg)
        if erreurs is not None and msg not in erreurs:
            erreurs.append(msg)
        return None

    site = Site.objects.create(code=site_code, nom=f"Site {site_code}")
    return site



def import_excel_log(
    file_path: str,
    sheet_name: str = 'RECEIVING-EXPORT Status',
    create_trajets: bool = True
) -> Dict[str, Any]:
    """
    Importe un fichier Excel au format LOG_TFZ-TOUBKAL dans la base Django.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    # Résolution des cellules fusionnées pour récupérer les valeurs sur chaque ligne
    merged_values = {}
    for rng in ws.merged_cells.ranges:
        top_val = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged_values[(r, c)] = top_val

    def get_cell(r: int, c: int) -> Any:
        return merged_values.get((r, c), ws.cell(r, c).value)

    # Cellules de quantité situées sous la première ligne d'une fusion verticale
    merged_qty_followers = {
        (r, c)
        for rng in ws.merged_cells.ranges
        for r in range(rng.min_row + 1, rng.max_row + 1)
        for c in range(rng.min_col, rng.max_col + 1)
        if c in (4, 9)
    }

    # Colonnes numériques (quantités et palettes) : une cellule fusionnée porte UNE
    # seule valeur. Elle n'est donc pas recopiée sur les lignes suivantes, sinon une
    # même quantité serait comptée plusieurs fois.
    QTY_COLS = (4, 5, 9, 10)

    def get_num(r: int, c: int) -> Any:
        return ws.cell(r, c).value if c in QTY_COLS else get_cell(r, c)

    def occurrence_is_new(key: tuple, filters: Dict[str, Any]) -> bool:
        """Idempotence par comptage d'occurrences.

        Deux lignes identiques (même date, site, PN, type et quantité) dans le fichier
        sont deux mouvements réels distincts. La n-ième occurrence d'une clé n'est créée
        que si la base en contient moins de n : un ré-import du même fichier (ou d'une
        version enrichie) ne crée donc aucun doublon, sans perdre de lignes réelles.
        """
        occurrences[key] = occurrences.get(key, 0) + 1
        return Mouvement.objects.filter(**filters).count() < occurrences[key]

    occurrences: Dict[tuple, int] = {}

    mouvements_crees = 0
    mouvements_doublons = 0
    receptions_crees = 0
    exports_crees = 0
    trajets_crees = 0
    erreurs: List[str] = []
    avertissements: List[str] = []
    lignes_ignorees_details: List[Dict[str, Any]] = []

    # En-têtes (lignes 1 à 3)
    if ws.max_row >= 1:
        lignes_ignorees_details.append({
            'ligne': 1,
            'categorie': 'en-tête',
            'raison': "En-tête : Titres de sections ('RECEIVING' / 'EXPORT')"
        })
    if ws.max_row >= 2:
        lignes_ignorees_details.append({
            'ligne': 2,
            'categorie': 'en-tête',
            'raison': "En-tête : Interligne sous les titres"
        })
    if ws.max_row >= 3:
        lignes_ignorees_details.append({
            'ligne': 3,
            'categorie': 'en-tête',
            'raison': "En-tête : Noms des colonnes (Date, Plant, PN, etc.)"
        })

    cur_date: Optional[datetime.date] = None
    cur_rec_plant: Optional[str] = None
    cur_exp_plant: Optional[str] = None
    cur_trajets: Optional[int] = None
    cur_remarque: Optional[str] = None

    for r in range(4, ws.max_row + 1):
        try:
            raw_cells = [ws.cell(r, c).value for c in range(1, 14)]

            # 1. Lignes complètement vides
            if all(v is None for v in raw_cells):
                lignes_ignorees_details.append({
                    'ligne': r,
                    'categorie': 'vide',
                    'raison': "Ligne vide"
                })
                continue

            d_raw = get_cell(r, 1)
            g_raw = get_cell(r, 7)

            # 2. Lignes Total Quantity intercalées
            if 'total quantity' in str(g_raw or '').lower() or 'total quantity' in str(d_raw or '').lower():
                lignes_ignorees_details.append({
                    'ligne': r,
                    'categorie': 'total',
                    'raison': "Ligne de total ('Total Quantity')"
                })
                continue

            parsed_d = parse_date_value(d_raw)
            if parsed_d and parsed_d != cur_date:
                cur_date = parsed_d
                cur_rec_plant = None
                cur_exp_plant = None
                cur_trajets = None
                cur_remarque = None

            # 3. Date introuvable
            if not cur_date:
                lignes_ignorees_details.append({
                    'ligne': r,
                    'categorie': 'donnée invalide',
                    'raison': f"Date manquante ou non reconnue ({repr(d_raw)})"
                })
                continue

            # Mise à jour des plants courants
            r_p_raw = get_cell(r, 2)
            if r_p_raw and str(r_p_raw).strip() != 'Without Reception':
                cur_rec_plant = normalize_plant_code(r_p_raw, cur_rec_plant)

            if g_raw:
                cur_exp_plant = normalize_plant_code(g_raw, cur_exp_plant)

            # Nombre de trajets et remarque éventuelle
            tr_val = parse_int_value(get_cell(r, 12))
            if tr_val is not None:
                cur_trajets = tr_val
            rem_val = get_cell(r, 13)
            if rem_val:
                cur_remarque = str(rem_val).strip()

            row_had_movement = False

            # --- Traitement RECEIVING ---
            r_pn = clean_pn_value(get_cell(r, 3))
            r_qty = parse_int_value(get_num(r, 4))
            if r_pn and r_qty is not None and r_qty > 0:
                site_code = cur_rec_plant or 'TFZ'
                site = get_or_create_site_with_warning(site_code, r, avertissements, erreurs)
                if not site:
                    lignes_ignorees_details.append({
                        'ligne': r,
                        'categorie': 'donnée invalide',
                        'raison': f"Donnée invalide : Le site '{site_code}' n'existe ni en base ni dans PLANT_MAPPING"
                    })
                    continue
                row_had_movement = True
                reference, _ = Reference.objects.get_or_create(code_pn=r_pn)
                nb_pal = parse_int_value(get_num(r, 5))

                # Idempotence par comptage d'occurrences (voir occurrence_is_new)
                rec_filters = dict(date_mouvement=cur_date, site=site, reference=reference,
                                   type_mouvement='RECEPTION', quantite=r_qty)
                if occurrence_is_new(('RECEPTION', cur_date, site.id, reference.id, r_qty), rec_filters):
                    Mouvement.objects.create(
                        date_mouvement=cur_date,
                        site=site,
                        reference=reference,
                        type_mouvement='RECEPTION',
                        quantite=r_qty,
                        nb_palettes=nb_pal,
                    )
                    mouvements_crees += 1
                    receptions_crees += 1
                else:
                    mouvements_doublons += 1

            # --- Traitement EXPORT ---
            e_pn = clean_pn_value(get_cell(r, 8))
            e_qty = parse_int_value(get_num(r, 9))
            if e_pn and e_qty is not None and e_qty > 0:
                site_code = cur_exp_plant or 'TFZ'
                site = get_or_create_site_with_warning(site_code, r, avertissements, erreurs)
                if not site:
                    lignes_ignorees_details.append({
                        'ligne': r,
                        'categorie': 'donnée invalide',
                        'raison': f"Donnée invalide : Le site '{site_code}' n'existe ni en base ni dans PLANT_MAPPING"
                    })
                    continue
                row_had_movement = True
                reference, _ = Reference.objects.get_or_create(code_pn=e_pn)
                nb_pal = parse_int_value(get_num(r, 10))

                # Idempotence par comptage d'occurrences (voir occurrence_is_new)
                exp_filters = dict(date_mouvement=cur_date, site=site, reference=reference,
                                   type_mouvement='EXPORT', quantite=e_qty)
                is_new = occurrence_is_new(('EXPORT', cur_date, site.id, reference.id, e_qty), exp_filters)

                mvt = None
                if is_new:
                    mvt = Mouvement.objects.create(
                        date_mouvement=cur_date,
                        site=site,
                        reference=reference,
                        type_mouvement='EXPORT',
                        quantite=e_qty,
                        nb_palettes=nb_pal,
                    )
                    mouvements_crees += 1
                    exports_crees += 1
                else:
                    # Occurrence déjà importée : elle a été rattachée à son trajet lors du
                    # premier import, on ne la rattache pas une seconde fois.
                    mouvements_doublons += 1

                # Création / liaison optionnelle du Trajet (nouveaux mouvements uniquement)
                if create_trajets and cur_trajets and mvt is not None:
                    trajet, t_created = Trajet.objects.get_or_create(
                        date_trajet=cur_date,
                        site=site,
                        remarque=cur_remarque or None,
                    )
                    if t_created:
                        trajets_crees += 1

                    TrajetMouvement.objects.get_or_create(
                        trajet=trajet,
                        mouvement=mvt,
                        defaults={
                            'quantite_transportee': e_qty,
                            'nb_palettes_transportees': nb_pal,
                        }
                    )

            if not row_had_movement:
                # Identification de la cause exacte
                raw_e_qty = get_num(r, 9)
                raw_r_qty = get_num(r, 4)
                if merged_qty_followers.intersection({(r, 4), (r, 9)}) and raw_e_qty is None and raw_r_qty is None:
                    lignes_ignorees_details.append({
                        'ligne': r,
                        'categorie': 'fusion',
                        'raison': "Suite d'une cellule de quantité fusionnée (quantité déjà comptée sur la première ligne)"
                    })
                elif raw_e_qty is not None and not e_pn:
                    lignes_ignorees_details.append({
                        'ligne': r,
                        'categorie': 'donnée invalide',
                        'raison': f"Donnée invalide : Quantité exportée ({raw_e_qty}) sans code PN"
                    })
                elif raw_r_qty is not None and not r_pn:
                    lignes_ignorees_details.append({
                        'ligne': r,
                        'categorie': 'donnée invalide',
                        'raison': f"Donnée invalide : Quantité reçue ({raw_r_qty}) sans code PN"
                    })
                elif e_pn and (e_qty is None or e_qty <= 0):
                    lignes_ignorees_details.append({
                        'ligne': r,
                        'categorie': 'donnée invalide',
                        'raison': f"Donnée invalide : Code PN export ({e_pn}) avec quantité invalide ({repr(raw_e_qty)})"
                    })
                elif r_pn and (r_qty is None or r_qty <= 0):
                    lignes_ignorees_details.append({
                        'ligne': r,
                        'categorie': 'donnée invalide',
                        'raison': f"Donnée invalide : Code PN réception ({r_pn}) avec quantité invalide ({repr(raw_r_qty)})"
                    })
                else:
                    lignes_ignorees_details.append({
                        'ligne': r,
                        'categorie': 'vide',
                        'raison': "Ligne sans mouvement exploitable"
                    })

        except Exception as e:
            err_msg = f"Ligne {r}: {str(e)}"
            logger.warning(err_msg)
            erreurs.append(err_msg)
            lignes_ignorees_details.append({
                'ligne': r,
                'categorie': 'donnée invalide',
                'raison': f"Donnée invalide (erreur : {str(e)})"
            })

    return {
        'mouvements_crees': mouvements_crees,
        'mouvements_doublons': mouvements_doublons,
        'receptions_crees': receptions_crees,
        'exports_crees': exports_crees,
        'trajets_crees': trajets_crees,
        'lignes_ignorees': len(lignes_ignorees_details),
        'lignes_ignorees_details': lignes_ignorees_details,
        'avertissements': avertissements,
        'erreurs': erreurs,
    }

