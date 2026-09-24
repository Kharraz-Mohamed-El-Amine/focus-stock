from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth.models import User
from datetime import date
from decimal import Decimal

from .models import (
    Site,
    Reference,
    Mouvement,
    BonLivraison,
    BonLivraisonMouvement,
    Trajet,
    TrajetMouvement,
    Facture,
    RapprochementFacture,
    Profil,
)


class StockViewsAndFormsTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        # Utilisateur superviseur par défaut pour les tests fonctionnels existants
        self.user_superviseur = User.objects.create_user(username='super_test', password='password123')
        Profil.objects.create(user=self.user_superviseur, role='SUPERVISEUR')
        self.client.force_login(self.user_superviseur)

        self.site_tfz = Site.objects.create(code='TFZ', nom='Tangier Free Zone')
        self.site_toubkal = Site.objects.create(code='TOUBKAL', nom='Site Toubkal')

        self.ref1 = Reference.objects.create(code_pn='PN-001', designation='Câble haute tension')
        self.ref2 = Reference.objects.create(code_pn='PN-002', designation='Connecteur standard')

        # Mouvements existants
        self.mvt_rec = Mouvement.objects.create(
            date_mouvement=date(2026, 9, 1),
            type_mouvement='RECEPTION',
            site=self.site_tfz,
            reference=self.ref1,
            quantite=500,
            nb_palettes=5,
        )
        self.mvt_exp1 = Mouvement.objects.create(
            date_mouvement=date(2026, 9, 2),
            type_mouvement='EXPORT',
            site=self.site_tfz,
            reference=self.ref1,
            quantite=200,
            nb_palettes=2,
        )
        self.mvt_exp2 = Mouvement.objects.create(
            date_mouvement=date(2026, 9, 3),
            type_mouvement='EXPORT',
            site=self.site_toubkal,
            reference=self.ref2,
            quantite=300,
            nb_palettes=3,
        )

    # ----------------------------------------------------
    # 1. Tests Mouvements
    # ----------------------------------------------------
    def test_mouvement_list_view(self):
        response = self.client.get(reverse('stock:mouvement_list'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'stock/mouvement_list.html')
        # Vérification du tri par date décroissante
        mouvements = list(response.context['page_obj'])
        self.assertGreaterEqual(mouvements[0].date_mouvement, mouvements[-1].date_mouvement)

    def test_mouvement_create_post_valid(self):
        payload = {
            'date_mouvement': '2026-09-05',
            'site': self.site_tfz.id,
            'type_mouvement': 'EXPORT',
            'reference': self.ref2.id,
            'quantite': 150,
            'nb_palettes': 1,
            'batch': 'BATCH-2026-X',
        }
        response = self.client.post(reverse('stock:mouvement_create'), data=payload, follow=True)
        self.assertRedirects(response, reverse('stock:mouvement_list'))
        self.assertTrue(Mouvement.objects.filter(batch='BATCH-2026-X').exists())
        self.assertContains(response, "Mouvement EXPORT pour PN-002 (150 pièces) enregistré avec succès.")

    def test_mouvement_create_post_invalid(self):
        payload = {
            'date_mouvement': '2026-09-05',
            'site': self.site_tfz.id,
            'type_mouvement': 'EXPORT',
            'reference': self.ref2.id,
            'quantite': -10,  # Quantité invalide
        }
        response = self.client.post(reverse('stock:mouvement_create'), data=payload)
        self.assertEqual(response.status_code, 200)
        self.assertFormError(response.context['form'], 'quantite', "La quantité doit être strictement positive.")

    # ----------------------------------------------------
    # 2. Tests Bon de Livraison
    # ----------------------------------------------------
    def test_bon_livraison_list_view(self):
        response = self.client.get(reverse('stock:bon_livraison_list'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'stock/bon_livraison_list.html')

    def test_bon_livraison_create_filter_eligible_movements(self):
        # Mouvement RECEPTION ne doit pas être proposé
        response = self.client.get(reverse('stock:bon_livraison_create'))
        self.assertEqual(response.status_code, 200)
        eligible = response.context['eligible_mouvements']
        self.assertNotIn(self.mvt_rec, eligible)
        self.assertIn(self.mvt_exp1, eligible)
        self.assertIn(self.mvt_exp2, eligible)

    def test_bon_livraison_create_post(self):
        payload = {
            'numero_bl': 'BL-2026-TEST',
            'date_bl': '2026-09-04',
            'site': self.site_tfz.id,
            'nb_palettes_total': 2,
            'poids_brut': '125.50',
            'mouvements': [self.mvt_exp1.id],
        }
        response = self.client.post(reverse('stock:bon_livraison_create'), data=payload, follow=True)
        self.assertRedirects(response, reverse('stock:bon_livraison_list'))

        bl = BonLivraison.objects.get(numero_bl='BL-2026-TEST')
        self.assertEqual(bl.site, self.site_tfz)
        self.assertTrue(BonLivraisonMouvement.objects.filter(bon_livraison=bl, mouvement=self.mvt_exp1).exists())
        self.assertContains(response, "Bon de livraison BL-2026-TEST (TFZ) enregistré avec succès.")

        # Vérifier que mvt_exp1 n'est plus éligible pour un autre BL
        response2 = self.client.get(reverse('stock:bon_livraison_create'))
        self.assertNotIn(self.mvt_exp1, response2.context['eligible_mouvements'])

    # ----------------------------------------------------
    # 3. Tests Trajet
    # ----------------------------------------------------
    def test_trajet_list_view(self):
        response = self.client.get(reverse('stock:trajet_list'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'stock/trajet_list.html')

    def test_trajet_create_post(self):
        payload = {
            'date_trajet': '2026-09-05',
            'site': self.site_tfz.id,
            'remarque': 'TAXI',
            'mouvements': [self.mvt_exp1.id],
            f'qty_{self.mvt_exp1.id}': '180',
            f'pal_{self.mvt_exp1.id}': '2',
        }
        response = self.client.post(reverse('stock:trajet_create'), data=payload, follow=True)
        self.assertRedirects(response, reverse('stock:trajet_list'))

        trajet = Trajet.objects.get(remarque='TAXI', site=self.site_tfz)
        tm = TrajetMouvement.objects.get(trajet=trajet, mouvement=self.mvt_exp1)
        self.assertEqual(tm.quantite_transportee, 180)
        self.assertEqual(tm.nb_palettes_transportees, 2)
        self.assertContains(response, "enregistré avec succès (1 mouvement(s) associé(s))")

    def test_trajet_create_refuse_depassement_quantite(self):
        """La somme des quantités transportées ne peut pas dépasser la quantité exportée."""
        premier = {
            'date_trajet': '2026-09-05', 'site': self.site_tfz.id, 'remarque': 'T1',
            'mouvements': [self.mvt_exp1.id], f'qty_{self.mvt_exp1.id}': '150',
        }
        self.client.post(reverse('stock:trajet_create'), data=premier)
        # Reste à transporter : 200 - 150 = 50 ; une demande de 60 doit être refusée
        second = {
            'date_trajet': '2026-09-06', 'site': self.site_tfz.id, 'remarque': 'T2',
            'mouvements': [self.mvt_exp1.id], f'qty_{self.mvt_exp1.id}': '60',
        }
        response = self.client.post(reverse('stock:trajet_create'), data=second)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "reste à transporter 50")
        self.assertFalse(Trajet.objects.filter(remarque='T2').exists())
        # Une demande de 50 (reste exact) est acceptée
        second[f'qty_{self.mvt_exp1.id}'] = '50'
        self.client.post(reverse('stock:trajet_create'), data=second)
        total = sum(tm.quantite_transportee for tm in self.mvt_exp1.trajet_mouvements.all())
        self.assertEqual(total, 200)

    # ----------------------------------------------------
    # 4. Tests Facture
    # ----------------------------------------------------
    def test_facture_list_view(self):
        response = self.client.get(reverse('stock:facture_list'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'stock/facture_list.html')

    def test_facture_create_post(self):
        payload = {
            'numero': 'FAC-2026-099',
            'date_facture': '2026-09-06',
            'site': self.site_tfz.id,
            'montant': '4500.00',
        }
        response = self.client.post(reverse('stock:facture_create'), data=payload, follow=True)
        self.assertRedirects(response, reverse('stock:facture_list'))

        self.assertTrue(Facture.objects.filter(numero='FAC-2026-099', montant=Decimal('4500.00')).exists())
        self.assertContains(response, "Facture FAC-2026-099 (TFZ) enregistrée avec succès.")

    # ----------------------------------------------------
    # 5. Tests Rapprochement Facture <-> Trajets
    # ----------------------------------------------------
    def test_rapprochement_detail_get_auto_checks_same_month(self):
        # Créer trajets en septembre 2026 et août 2026
        t_sep = Trajet.objects.create(date_trajet=date(2026, 9, 10), site=self.site_tfz)
        t_aout = Trajet.objects.create(date_trajet=date(2026, 8, 20), site=self.site_tfz)

        facture = Facture.objects.create(
            numero='FAC-SEP-01',
            date_facture=date(2026, 9, 30),
            site=self.site_tfz,
            montant=Decimal('1000.00'),
            nb_trajets_factures=1
        )

        response = self.client.get(reverse('stock:rapprochement_detail', kwargs={'facture_id': facture.id}))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'stock/rapprochement_detail.html')

        eligible = response.context['eligible_trajets']
        self.assertIn(t_sep, eligible)
        self.assertIn(t_aout, eligible)

        # Vérifier que le trajet de septembre est coché par défaut, mais pas celui d'août
        t_sep_obj = next(t for t in eligible if t.id == t_sep.id)
        t_aout_obj = next(t for t in eligible if t.id == t_aout.id)
        self.assertTrue(t_sep_obj.is_checked)
        self.assertFalse(t_aout_obj.is_checked)

    def test_rapprochement_excludes_trajets_linked_to_other_facture(self):
        t1 = Trajet.objects.create(date_trajet=date(2026, 9, 10), site=self.site_tfz)
        t2 = Trajet.objects.create(date_trajet=date(2026, 9, 11), site=self.site_tfz)

        # Facture 1 liée à Trajet 1
        f1 = Facture.objects.create(numero='FAC-1', date_facture=date(2026, 9, 30), site=self.site_tfz, nb_trajets_factures=1)
        RapprochementFacture.objects.create(facture=f1, trajet=t1, statut='OK')

        # Facture 2 (même site)
        f2 = Facture.objects.create(numero='FAC-2', date_facture=date(2026, 9, 30), site=self.site_tfz, nb_trajets_factures=1)

        response = self.client.get(reverse('stock:rapprochement_detail', kwargs={'facture_id': f2.id}))
        self.assertEqual(response.status_code, 200)

        eligible = response.context['eligible_trajets']
        # t1 doit être exclu car déjà lié à f1
        self.assertNotIn(t1, eligible)
        self.assertIn(t2, eligible)

    def test_rapprochement_calcul_ecart_ok(self):
        t1 = Trajet.objects.create(date_trajet=date(2026, 9, 10), site=self.site_tfz)
        t2 = Trajet.objects.create(date_trajet=date(2026, 9, 11), site=self.site_tfz)

        facture = Facture.objects.create(
            numero='FAC-OK',
            date_facture=date(2026, 9, 30),
            site=self.site_tfz,
            nb_trajets_factures=2
        )

        payload = {
            'trajets': [t1.id, t2.id]
        }
        response = self.client.post(reverse('stock:rapprochement_detail', kwargs={'facture_id': facture.id}), data=payload)
        self.assertEqual(response.status_code, 200)

        # 2 déclarés vs 2 réels -> statut OK
        rapps = RapprochementFacture.objects.filter(facture=facture)
        self.assertEqual(rapps.count(), 2)
        self.assertTrue(all(r.statut == 'OK' for r in rapps))
        self.assertEqual(facture.statut_rapprochement, 'OK')
        self.assertEqual(response.context['reconciliation_result']['statut'], 'OK')

    def test_rapprochement_calcul_ecart_manquant(self):
        t1 = Trajet.objects.create(date_trajet=date(2026, 9, 10), site=self.site_tfz)

        facture = Facture.objects.create(
            numero='FAC-MANQUANT',
            date_facture=date(2026, 9, 30),
            site=self.site_tfz,
            nb_trajets_factures=3  # 3 facturés
        )

        payload = {
            'trajets': [t1.id]  # seulement 1 réel
        }
        response = self.client.post(reverse('stock:rapprochement_detail', kwargs={'facture_id': facture.id}), data=payload)
        self.assertEqual(response.status_code, 200)

        rapps = RapprochementFacture.objects.filter(facture=facture)
        self.assertEqual(rapps.count(), 1)
        self.assertEqual(rapps.first().statut, 'ECART')
        self.assertEqual(facture.statut_rapprochement, 'ECART')
        self.assertEqual(response.context['reconciliation_result']['statut'], 'ECART')
        self.assertContains(response, "ECART - 2 trajets manquants")

    def test_rapprochement_calcul_ecart_en_trop(self):
        t1 = Trajet.objects.create(date_trajet=date(2026, 9, 10), site=self.site_tfz)
        t2 = Trajet.objects.create(date_trajet=date(2026, 9, 11), site=self.site_tfz)

        facture = Facture.objects.create(
            numero='FAC-TROP',
            date_facture=date(2026, 9, 30),
            site=self.site_tfz,
            nb_trajets_factures=1  # 1 facturé
        )

        payload = {
            'trajets': [t1.id, t2.id]  # 2 réels
        }
        response = self.client.post(reverse('stock:rapprochement_detail', kwargs={'facture_id': facture.id}), data=payload)
        self.assertEqual(response.status_code, 200)

        rapps = RapprochementFacture.objects.filter(facture=facture)
        self.assertEqual(rapps.count(), 2)
        self.assertTrue(all(r.statut == 'ECART' for r in rapps))
        self.assertEqual(facture.statut_rapprochement, 'ECART')
        self.assertEqual(response.context['reconciliation_result']['statut'], 'ECART')
        self.assertContains(response, "ECART - 1 trajet en trop")

    def test_rapprochement_statut_updates_correctly(self):
        t1 = Trajet.objects.create(date_trajet=date(2026, 9, 10), site=self.site_tfz)
        t2 = Trajet.objects.create(date_trajet=date(2026, 9, 11), site=self.site_tfz)

        facture = Facture.objects.create(
            numero='FAC-UPDATE',
            date_facture=date(2026, 9, 30),
            site=self.site_tfz,
            nb_trajets_factures=2
        )

        # 1er POST : avec 1 seul trajet (écart)
        self.client.post(reverse('stock:rapprochement_detail', kwargs={'facture_id': facture.id}), data={'trajets': [t1.id]})
        self.assertEqual(RapprochementFacture.objects.get(facture=facture, trajet=t1).statut, 'ECART')
        self.assertEqual(facture.statut_rapprochement, 'ECART')

        # 2ème POST : correction avec les 2 trajets (conforme)
        self.client.post(reverse('stock:rapprochement_detail', kwargs={'facture_id': facture.id}), data={'trajets': [t1.id, t2.id]})
        rapps = RapprochementFacture.objects.filter(facture=facture)
        self.assertEqual(rapps.count(), 2)
        self.assertTrue(all(r.statut == 'OK' for r in rapps))
        self.assertEqual(facture.statut_rapprochement, 'OK')

    def test_facture_list_status_badges_and_actions(self):
        f_non_rap = Facture.objects.create(numero='FAC-NONE', date_facture=date(2026, 9, 1), site=self.site_tfz)
        f_ok = Facture.objects.create(numero='FAC-OK2', date_facture=date(2026, 9, 2), site=self.site_tfz, nb_trajets_factures=1)
        t_ok = Trajet.objects.create(date_trajet=date(2026, 9, 2), site=self.site_tfz)
        RapprochementFacture.objects.create(facture=f_ok, trajet=t_ok, statut='OK')

        response = self.client.get(reverse('stock:facture_list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Non rapproché")
        self.assertContains(response, "OK")
        self.assertContains(response, reverse('stock:rapprochement_detail', kwargs={'facture_id': f_ok.id}))

    def test_trajets_non_factures_detection(self):
        # Trajet antérieur à la facture non rattaché
        t_orphelin = Trajet.objects.create(date_trajet=date(2026, 8, 15), site=self.site_tfz)
        # Trajet postérieur à la facture
        t_post = Trajet.objects.create(date_trajet=date(2026, 10, 5), site=self.site_tfz)

        facture = Facture.objects.create(
            numero='FAC-ORPHELINS',
            date_facture=date(2026, 9, 30),
            site=self.site_tfz
        )

        response = self.client.get(reverse('stock:rapprochement_detail', kwargs={'facture_id': facture.id}))
        self.assertEqual(response.status_code, 200)
        trajets_non_factures = response.context['trajets_non_factures']
        self.assertIn(t_orphelin, trajets_non_factures)
        self.assertNotIn(t_post, trajets_non_factures)

    # ----------------------------------------------------
    # 6. Tests Dashboard (Vue Home)
    # ----------------------------------------------------
    def test_home_dashboard_stock_table(self):
        # self.mvt_rec : RECEPTION, TFZ, ref1, 500
        # self.mvt_exp1 : EXPORT, TFZ, ref1, 200
        # Stock ref1 / TFZ = 500 - 200 = 300
        response = self.client.get(reverse('stock:home'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'stock/home.html')

        stock_data = response.context['stock_data']
        item_ref1_tfz = next(
            (item for item in stock_data if item['reference__code_pn'] == self.ref1.code_pn and item['site__code'] == 'TFZ'),
            None
        )
        self.assertIsNotNone(item_ref1_tfz)
        self.assertEqual(item_ref1_tfz['total_receptions'], 500)
        self.assertEqual(item_ref1_tfz['total_exports'], 200)
        self.assertEqual(item_ref1_tfz['stock_actuel'], 300)
        self.assertContains(response, "+300")

    def test_home_dashboard_kpis_and_orphan_trips(self):
        today = date.today()
        # Créer un trajet ce mois pour TFZ
        t_ce_mois = Trajet.objects.create(date_trajet=today, site=self.site_tfz)
        # Créer un trajet orphelin (non lié)
        t_orphelin = Trajet.objects.create(date_trajet=date(2026, 8, 1), site=self.site_toubkal)

        # Créer une facture en écart
        f_ecart = Facture.objects.create(
            numero='FAC-ECART-DASH',
            date_facture=today,
            site=self.site_tfz,
            nb_trajets_factures=5
        )
        # Lier seulement 1 trajet -> écart
        RapprochementFacture.objects.create(facture=f_ecart, trajet=t_ce_mois, statut='ECART')

        response = self.client.get(reverse('stock:home'))
        self.assertEqual(response.status_code, 200)

        # Vérifier factures en écart
        self.assertGreaterEqual(response.context['nb_factures_ecart'], 1)

        # Vérifier trajets orphelins
        self.assertGreaterEqual(response.context['total_trajets_orphelins'], 1)
        self.assertContains(response, reverse('stock:trajet_list') + "?orphelin=1")

        # Vérifier les 10 derniers mouvements
        derniers_mvts = response.context['derniers_mouvements']
        self.assertLessEqual(len(derniers_mvts), 10)

    def test_trajet_list_filter_orphelin(self):
        t_lie = Trajet.objects.create(date_trajet=date(2026, 9, 1), site=self.site_tfz)
        t_orphelin = Trajet.objects.create(date_trajet=date(2026, 9, 2), site=self.site_tfz)

        f = Facture.objects.create(numero='FAC-LIEE', date_facture=date(2026, 9, 3), site=self.site_tfz)
        RapprochementFacture.objects.create(facture=f, trajet=t_lie, statut='OK')

        # Test sans filtre
        res_all = self.client.get(reverse('stock:trajet_list'))
        self.assertContains(res_all, t_lie.date_trajet.strftime('%d/%m/%Y'))

        # Test avec filtre orphelin
        res_orphelin = self.client.get(reverse('stock:trajet_list') + '?orphelin=1')
        self.assertEqual(res_orphelin.status_code, 200)
        trajets_result = list(res_orphelin.context['page_obj'])
        self.assertIn(t_orphelin, trajets_result)
        self.assertNotIn(t_lie, trajets_result)


class AuthenticationAndRoleTestCase(TestCase):
    def setUp(self):
        self.site = Site.objects.create(code='TFZ', nom='Tangier Free Zone')
        self.ref = Reference.objects.create(code_pn='PN-TEST', designation='Composant Test')
        self.facture = Facture.objects.create(
            numero='FAC-AUTH-TEST',
            date_facture=date(2026, 9, 10),
            site=self.site
        )

        # 1. Superviseur
        self.superviseur = User.objects.create_user(
            username='super_user',
            password='Password123!',
            first_name='Samir'
        )
        Profil.objects.create(user=self.superviseur, role='SUPERVISEUR')

        # 2. Chef d'équipe
        self.chef_equipe = User.objects.create_user(
            username='chef_user',
            password='Password123!',
            first_name='Karim'
        )
        Profil.objects.create(user=self.chef_equipe, role='CHEF_EQUIPE')

        # Client non connecté
        self.anon_client = Client()

    def test_anonymous_user_redirected_to_login(self):
        urls = [
            reverse('stock:home'),
            reverse('stock:mouvement_list'),
            reverse('stock:mouvement_create'),
            reverse('stock:bon_livraison_list'),
            reverse('stock:bon_livraison_create'),
            reverse('stock:trajet_list'),
            reverse('stock:trajet_create'),
            reverse('stock:facture_list'),
            reverse('stock:facture_create'),
            reverse('stock:rapprochement_detail', kwargs={'facture_id': self.facture.id}),
        ]
        for url in urls:
            with self.subTest(url=url):
                res = self.anon_client.get(url)
                self.assertEqual(res.status_code, 302)
                self.assertIn('/login/', res.url)

    def test_chef_equipe_can_consult_lists_and_create_mouvement(self):
        client = Client()
        client.force_login(self.chef_equipe)

        # Consultation autorisée
        self.assertEqual(client.get(reverse('stock:home')).status_code, 200)
        self.assertEqual(client.get(reverse('stock:mouvement_list')).status_code, 200)
        self.assertEqual(client.get(reverse('stock:bon_livraison_list')).status_code, 200)
        self.assertEqual(client.get(reverse('stock:trajet_list')).status_code, 200)
        self.assertEqual(client.get(reverse('stock:facture_list')).status_code, 200)

        # Création de mouvement autorisée
        self.assertEqual(client.get(reverse('stock:mouvement_create')).status_code, 200)
        post_res = client.post(reverse('stock:mouvement_create'), {
            'date_mouvement': '2026-09-12',
            'type_mouvement': 'RECEPTION',
            'site': self.site.id,
            'reference': self.ref.id,
            'quantite': 150,
            'nb_palettes': 1,
        })
        self.assertRedirects(post_res, reverse('stock:mouvement_list'))
        self.assertTrue(Mouvement.objects.filter(quantite=150, reference=self.ref).exists())

    def test_chef_equipe_forbidden_from_restricted_actions(self):
        client = Client()
        client.force_login(self.chef_equipe)

        restricted_urls = [
            reverse('stock:bon_livraison_create'),
            reverse('stock:trajet_create'),
            reverse('stock:facture_create'),
            reverse('stock:rapprochement_detail', kwargs={'facture_id': self.facture.id}),
        ]

        for url in restricted_urls:
            with self.subTest(url=url):
                # GET interdit
                res_get = client.get(url)
                self.assertEqual(res_get.status_code, 403)
                # POST interdit
                res_post = client.post(url, {})
                self.assertEqual(res_post.status_code, 403)

    def test_superviseur_has_full_access(self):
        client = Client()
        client.force_login(self.superviseur)

        # Accès complet à la création et au rapprochement
        self.assertEqual(client.get(reverse('stock:bon_livraison_create')).status_code, 200)
        self.assertEqual(client.get(reverse('stock:trajet_create')).status_code, 200)
        self.assertEqual(client.get(reverse('stock:facture_create')).status_code, 200)
        self.assertEqual(
            client.get(reverse('stock:rapprochement_detail', kwargs={'facture_id': self.facture.id})).status_code,
            200
        )

    def test_login_page_renders_and_authenticates(self):
        client = Client()
        # Rendu formulaire login
        res_get = client.get(reverse('login'))
        self.assertEqual(res_get.status_code, 200)
        self.assertTemplateUsed(res_get, 'registration/login.html')

        # Authentification réussie
        res_post = client.post(reverse('login'), {
            'username': 'super_user',
            'password': 'Password123!'
        })
        self.assertRedirects(res_post, reverse('stock:home'))

        # Mauvais mot de passe
        res_fail = client.post(reverse('login'), {
            'username': 'super_user',
            'password': 'WrongPassword!'
        })
        self.assertEqual(res_fail.status_code, 200)
        self.assertContains(res_fail, "Identifiants invalides")

    def test_logout_view(self):
        client = Client()
        client.force_login(self.superviseur)

        # Déconnexion
        res_logout = client.post(reverse('logout'))
        self.assertRedirects(res_logout, reverse('login'))

        # Vérifier que le client n'est plus authentifié
        res_home = client.get(reverse('stock:home'))
        self.assertEqual(res_home.status_code, 302)





class ExcelImportTestCase(TestCase):
    """Import Excel : lignes identiques réelles, cellules de quantité fusionnées, ré-import."""

    def setUp(self):
        import os
        import tempfile
        import openpyxl
        Site.objects.create(code='TFZ', nom='TE TFZ')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'RECEIVING-EXPORT Status'
        ws['B1'], ws['G1'] = 'RECEIVING', 'EXPORT'
        ws.append([])
        ws.append(['Date', 'Plant', 'PN', 'Qty', 'Pal', None, 'Plant', 'PN', 'Qty', 'Pal', None, 'Trajets', None])
        # Ligne 4-5 : deux exports réels identiques (même PN, même quantité)
        ws.append([date(2026, 8, 3), 'Without Reception', None, None, None, None, 'TFZ', '111-1', 2400, 1, None, 1, None])
        ws.append([None, None, None, None, None, None, None, '111-1', 2400, 1, None, None, None])
        # Lignes 6-8 : UNE quantité fusionnée sur trois lignes (I6:I8) pour un même PN
        ws.append([None, None, None, None, None, None, None, '222-1', 21600, None, None, None, None])
        ws.append([None, None, None, None, None, None, None, None, None, None, None, None, None])
        ws.append([None, None, None, None, None, None, None, None, None, None, None, None, None])
        ws.merge_cells('A4:A8')
        ws.merge_cells('G4:G8')
        ws.merge_cells('L4:L8')
        ws.merge_cells('H6:H8')
        ws.merge_cells('I6:I8')
        fd, self.path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(self.path)

    def tearDown(self):
        import os
        os.remove(self.path)

    def test_lignes_identiques_conservees_et_fusion_comptee_une_fois(self):
        from stock.importers.excel_import import import_excel_log
        res = import_excel_log(self.path)
        self.assertEqual(res['exports_crees'], 3)
        self.assertEqual(Mouvement.objects.filter(reference__code_pn='111-1').count(), 2)
        self.assertEqual(Mouvement.objects.filter(reference__code_pn='222-1').count(), 1)
        self.assertEqual(sum(Mouvement.objects.values_list('quantite', flat=True)), 2400 * 2 + 21600)

    def test_reimport_idempotent(self):
        from stock.importers.excel_import import import_excel_log
        import_excel_log(self.path)
        res = import_excel_log(self.path)
        self.assertEqual(res['mouvements_crees'], 0)
        self.assertEqual(res['mouvements_doublons'], 3)
        self.assertEqual(Mouvement.objects.count(), 3)
        # Aucun double rattachement au trajet lors du ré-import
        self.assertEqual(TrajetMouvement.objects.count(), 3)
